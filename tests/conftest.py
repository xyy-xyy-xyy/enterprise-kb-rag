"""共享 fixture。

两条铁律：
1. **不依赖 `data/docs/` 下的真实文档** —— 那些文件全被 .gitignore 排除，
   CI 里根本不存在。所有语料一律用 tmp_path 动态生成（pymupdf / python-docx）。
2. **不碰真实密钥文件** —— 需要密钥时 monkeypatch `config.SM4_KEY`，
   走"环境变量优先"那条分支，绝不触发自动生成、不写 `data/.sm4_key`。
"""

import io
import os
import sys

# 必须在 import gradio 之前设：否则导入 app.main 时 Gradio 会去
# https://api.gradio.app/pkg-version 查版本。CI 上这是外网请求，会拖慢甚至挂住
# 整个测试进程；本机则是白白发一次请求。app/main.py 在模块作用域调 build_ui()，
# 所以只要 import app.main 就必然触发。
os.environ.setdefault("GRADIO_ANALYTICS_ENABLED", "False")

import docx  # noqa: E402
import pymupdf  # noqa: E402
import pytest  # noqa: E402
from langchain_core.documents import Document  # noqa: E402

# 让测试能 `import app.*`：pytest 默认只把 tests/ 加进 sys.path，
# `python -m pytest` 恰好能用是因为 cwd 在 path 里，裸 `pytest`（CI 的用法）就不行。
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)


@pytest.fixture(scope="session", autouse=True)
def _warm_up_jieba():
    """预热 jieba 前缀词典。

    首次调用会打印 "Building prefix dict…" / "Loading model…" 等噪音并耗时 1-2 秒。
    预热一次，避免它混进第一个测试的输出里被误读成失败。
    """
    import jieba

    jieba.lcut("预热分词")
    return True


# --------------------------- PDF 语料 ---------------------------


def _build_pdf_bytes() -> bytes:
    """3 页 PDF：第 1 页标题+正文，第 2 页**空白**，第 3 页标题+正文。

    空白页是故意的 —— 它必须不产生任何分块，否则页码会被后续内容顶偏。
    中文必须用 pymupdf 内置 CJK 字体 `china-s`，默认的 helv 写不出汉字。
    """
    doc = pymupdf.open()

    page1 = doc.new_page()
    page1.insert_text((72, 90), "第一章 住宿标准", fontname="china-s", fontsize=15)
    page1.insert_text(
        (72, 120), "一线城市住宿费上限为每晚 800 元。", fontname="china-s", fontsize=11
    )

    doc.new_page()  # 第 2 页：空白

    page3 = doc.new_page()
    page3.insert_text((72, 90), "第二章 报销时限", fontname="china-s", fontsize=15)
    page3.insert_text(
        (72, 120), "出差结束后 15 个工作日内提交报销单据。", fontname="china-s", fontsize=11
    )

    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    return _build_pdf_bytes()


@pytest.fixture
def sample_pdf_path(tmp_path, sample_pdf_bytes) -> str:
    path = tmp_path / "sample.pdf"
    path.write_bytes(sample_pdf_bytes)
    return str(path)


# --------------------------- Word 语料 ---------------------------


def _build_docx_bytes() -> bytes:
    """含标题段落、正文段落、**空段落**、标题层级与一个表格的 docx。

    空段落的存在是为了验证"空段落占号"：段落号 = paragraphs 下标 + 1，
    空段落不产生内容块但**照样占一个号**，用户在 Word 里数段落才能对上。
    """
    document = docx.Document()

    document.add_heading("员工考勤与请假管理办法", level=1)  # 段落号 1（Heading 样式 → 标题）
    document.add_paragraph("本办法自 2026 年 1 月 1 日起施行。")  # 段落号 2
    document.add_paragraph("")  # 段落号 3：空段落，占号但不产生块
    document.add_heading("第一章 年假", level=2)  # 段落号 4（章节号开头 → 标题）
    document.add_paragraph("入职满一年可享受带薪年假 5 天。")  # 段落号 5

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "标准"
    table.cell(1, 0).text = "年假"
    table.cell(1, 1).text = "5 天"

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    return _build_docx_bytes()


@pytest.fixture
def sample_docx_path(tmp_path, sample_docx_bytes) -> str:
    path = tmp_path / "sample.docx"
    path.write_bytes(sample_docx_bytes)
    return str(path)


# --------------------------- 检索层假语料 ---------------------------


@pytest.fixture
def fake_documents() -> list[Document]:
    """4 个分块，覆盖 PDF 与 Word 两种定位形态。

    A/C 的正文含"住宿费"，B 含"年假"，D 是"住宿费"与"年假"都不含的干扰项，
    供 BM25 的 0 分过滤用例使用。
    """
    return [
        Document(
            page_content="一线城市住宿费上限为每晚 800 元。",
            metadata={
                "source": "C:\\docs\\星尘计划差旅报销管理办法.docx",
                "file_name": "星尘计划差旅报销管理办法.docx",
                "page": 0,
                "chunk_id": "星尘计划差旅报销管理办法.docx_p0_0",
                "paragraph_start": 2,
                "paragraph_end": 4,
            },
        ),
        Document(
            page_content="入职满一年可享受带薪年假 5 天。",
            metadata={
                "source": "C:\\docs\\员工考勤与请假管理办法.docx",
                "file_name": "员工考勤与请假管理办法.docx",
                "page": 0,
                "chunk_id": "员工考勤与请假管理办法.docx_p0_0",
                "paragraph_start": 5,
                "paragraph_end": 5,
            },
        ),
        Document(
            page_content="住宿费超标部分需由员工自行承担。",
            metadata={
                "source": "C:\\docs\\外派人员费用管理规定.pdf",
                "file_name": "外派人员费用管理规定.pdf",
                "page": 3,
                "chunk_id": "外派人员费用管理规定.pdf_p3_0",
                "paragraph_start": 2,
                "paragraph_end": 4,
            },
        ),
        Document(
            page_content="本系统支持多轮对话与流式输出。",
            metadata={
                "source": "C:\\docs\\README.md",
                "file_name": "README.md",
                "page": 1,
                "chunk_id": "README.md_p1_0",
                "paragraph_start": 1,
                "paragraph_end": 1,
            },
        ),
    ]


@pytest.fixture
def sample_sm4_key() -> str:
    """固定的测试密钥（32 位 hex = 16 字节），保证测试可复现。"""
    return "0123456789abcdeffedcba9876543210"
